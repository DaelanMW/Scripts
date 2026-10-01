#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Conversor Universal de coordenadas - SAN JUAN, ARGENTINA
========================================================
Procesa todos los .xlsx / .xlsm de la carpeta donde está este script, detecta las
coordenadas de cada hoja y agrega columnas con:
    WGS84 decimal | POSGAR 2007 (faja 2 o 3) | UTM 19S | grados-minutos-segundos

Guarda el resultado en un archivo NUEVO  "<nombre>_convertido.xlsx"  (no pisa el original).

Como todo está en San Juan, el sistema se deduce por los rangos de los valores:
    Latitud   : 28° a 33° Sur            (decimal, GMS, con o sin signo)
    Longitud  : 66° a 71° Oeste
    POSGAR    : Este 2,2 a 3,7 millones  |  Norte 6,2 a 7,0 millones
    UTM 19S   : Este 250 mil a 850 mil   |  Norte 6,2 a 7,0 millones
Como los rangos no se superponen, NO importa el orden de las columnas
(lat/lon, lon/lat, este/norte, norte/este) ni si faltan los signos.

Reconoce: decimales (punto o coma), GMS en una celda (31°38'12.5"S, S 31 38 12.5,
31d38m12s, 31-38-12.5, 313812.5S, 31°38.208'), GMS en 3 columnas (grados|min|seg),
par de coordenadas en una sola celda ("-31.5, -69.4", "X: .. Y: .."), con o sin
encabezados, en cualquier fila de la hoja y en todas las hojas del libro.

Dependencias:  pip install openpyxl pyproj
"""
import math
import re
import subprocess
import sys
from collections import Counter
from pathlib import Path


def _instalar():
    import importlib
    faltan = []
    for mod, pkg in (("openpyxl", "openpyxl"), ("pyproj", "pyproj")):
        try:
            importlib.import_module(mod)
        except ImportError:
            faltan.append(pkg)
    if faltan:
        print("Instalando:", ", ".join(faltan))
        if subprocess.call([sys.executable, "-m", "pip", "install", *faltan]) != 0:
            sys.exit(f'No se pudo instalar. Ejecutá:\n    "{sys.executable}" -m pip install {" ".join(faltan)}')


_instalar()

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pyproj import Transformer

# ======================================================================
# CONSTANTES DE SAN JUAN
# ======================================================================
LAT_ABS = (27.5, 33.6)          # |latitud|  (°S)
LON_ABS = (65.8, 71.0)          # |longitud| (°O)
POS_E = (2_200_000, 3_700_000)  # POSGAR 2007 fajas 2 y 3
UTM_E = (250_000, 850_000)      # UTM huso 19S
NORTE = (6_200_000, 7_000_000)  # vale para POSGAR y UTM sur
EPSG_POSGAR = {2: 5344, 3: 5345}
EPSG_UTM19S = 32719

T_A_POS = {f: Transformer.from_crs(4326, e, always_xy=True) for f, e in EPSG_POSGAR.items()}
T_DE_POS = {f: Transformer.from_crs(e, 4326, always_xy=True) for f, e in EPSG_POSGAR.items()}
T_A_UTM = Transformer.from_crs(4326, EPSG_UTM19S, always_xy=True)
T_DE_UTM = Transformer.from_crs(EPSG_UTM19S, 4326, always_xy=True)

VACIOS = {"", "nan", "none", "null", "nd", "s/d", "sd", "n/a", "na", "-", "--", "#n/a", "#value!", "s/n"}


def faja_de_lon(lon):
    """San Juan usa fajas 2 (hasta 67°30'O) y 3."""
    return 2 if lon < -67.5 else 3


# ======================================================================
# LECTURA DE VALORES
# ======================================================================
def es_vacio(v):
    if v is None:
        return True
    if isinstance(v, float) and math.isnan(v):
        return True
    return isinstance(v, str) and v.strip().lower() in VACIOS


_RE_NUM = re.compile(r"^[-+]?(\d+\.?\d*|\.\d+)([eE][-+]?\d+)?$")


def parse_numero(v):
    if es_vacio(v) or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v) if math.isfinite(v) else None
    if not isinstance(v, str):
        return None
    s = v.strip().replace("\u00a0", " ").replace("\u202f", " ")
    s = s.replace("−", "-").replace("–", "-").replace("—", "-")
    s = re.sub(r"\s*(mts?|metros?|m)\.?$", "", s, flags=re.I)
    if re.fullmatch(r"[-+]?\d{1,3}( \d{3})+([.,]\d+)?", s):
        s = s.replace(" ", "")
    if " " in s:
        return None
    if "," in s and "." in s:
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif "," in s:
        if re.fullmatch(r"[-+]?\d{1,3}(,\d{3}){2,}", s):
            s = s.replace(",", "")
        else:
            s = s.replace(",", ".")
    elif s.count(".") > 1:
        if re.fullmatch(r"[-+]?\d{1,3}(\.\d{3})+", s):
            s = s.replace(".", "")
        else:
            return None
    if _RE_NUM.match(s):
        x = float(s)
        return x if math.isfinite(x) else None
    return None


def _compacto(x):
    """DDMMSS(.s) / DDDMMSS(.s) / DDMM(.m) escrito como un único número."""
    a = abs(x)
    ent = int(a)
    frac = a - ent
    n = len(str(ent))
    if n in (6, 7):
        ss, mm, dd = ent % 100 + frac, (ent // 100) % 100, ent // 10000
    elif n in (4, 5):
        mm, ss, dd = ent % 100 + frac, 0, ent // 100
    else:
        return None
    if mm >= 60 or ss >= 60 or dd > 180:
        return None
    return dd + mm / 60 + ss / 3600


def _armar(partes, hem):
    if hem:
        hem = {"O": "W"}.get(hem.upper(), hem.upper())
    if not 1 <= len(partes) <= 3:
        return None, None
    try:
        nums = [float(p) for p in partes]
    except ValueError:
        return None, None
    neg = partes[0].startswith("-") or math.copysign(1, nums[0]) < 0 or hem in ("S", "W")
    d = abs(nums[0])
    if len(nums) == 1:
        if d >= 1000:
            val = _compacto(d)
            if val is None:
                return None, None
        elif d > 180:
            return None, None
        else:
            val = d
    else:
        if d != int(d):
            return None, None
        m = nums[1]
        if m < 0 or m >= 60:
            return None, None
        val = d + m / 60
        if len(nums) == 3:
            s = nums[2]
            if m != int(m) or s < 0 or s >= 60:
                return None, None
            val += s / 3600
    return (-val if neg else val), hem


_PALABRAS_HEM = [(r"\bsur\b|\bsouth\b", "S"), (r"\bnorte\b|\bnorth\b", "N"),
                 (r"\boeste\b|\bwest\b", "W"), (r"\beste\b|\beast\b", "E")]
_RE_LETRAS_DMS = re.compile(
    r"\s*([NSEWOnsewo])?[\s:=]*([-+]?\d+(?:[.,]\d+)?)\s*[dD]\s*"
    r"(?:(\d+(?:[.,]\d+)?)\s*[mM]\s*(?:(\d+(?:[.,]\d+)?)\s*[sS]?)?)?\s*([NSEWO])?\s*")


def parse_angulo_texto(s):
    """Texto en GMS (o decimal con letras) -> (grados decimales, hemisferio|None)."""
    t = s.strip()
    if not t:
        return None, None
    t = t.replace("−", "-").replace("–", "-").replace("—", "-")
    t = re.sub(r"[º˚°]", "°", t)
    t = re.sub(r"[′’´`‘]", "'", t)
    t = re.sub(r"[″”“]", '"', t)
    t = t.replace("''", '"')
    hem = None
    for pat, h in _PALABRAS_HEM:
        if re.search(pat, t, re.I):
            hem = h
            t = re.sub(pat, " ", t, flags=re.I)
    m = _RE_LETRAS_DMS.fullmatch(t)
    if m:
        partes = [p.replace(",", ".") for p in (m.group(2), m.group(3), m.group(4)) if p]
        return _armar(partes, m.group(1) or m.group(5) or hem)
    m = re.match(r"^\s*([NSEWOnsewo])[\s:=.]*(?=[-+]?\d)", t)
    if m:
        hem = hem or m.group(1).upper()
        t = t[m.end():]
    else:
        m = re.search(r"(?<=[\d'\"°.\s])\s*([NSEWOnsewo])\s*$", t)
        if m:
            hem = hem or m.group(1).upper()
            t = t[:m.start()]
    t = re.sub(r"\b(grados?|degrees?|deg|minutos?|mins?|minutes?|segundos?|segs?|secs?|seconds?|"
               r"latitud|latitude|lat|longitud|longitude|long|lon|lng|coord\w*)\b", " ", t, flags=re.I)
    if re.search(r"[A-Za-z]", t):
        return None, None
    t = re.sub(r"[°'\"]", " ", t).replace(":", " ")
    t = re.sub(r"(?<=\d)\s*-\s*(?=\d)", " ", t)
    if t.count(",") >= 2 and "." not in t:
        t = t.replace(",", " ")
    else:
        t = t.replace(",", ".")
    patron = r"[-+]?\d+(?:\.\d+)?"
    partes = re.findall(patron, t)
    if re.sub(patron, "", t).strip(" \t.;"):
        return None, None
    return _armar(partes, hem)


def _dividir_par(t):
    partes = [p for p in re.split(r"\s*[;|\t\n]\s*|\s+/\s+|\s+y\s+", t) if p.strip()]
    if len(partes) == 2:
        return partes
    partes = [p for p in re.split(r",\s+", t) if p.strip()]
    if len(partes) == 2:
        return partes
    if t.count(",") == 1 and re.fullmatch(r"\s*\S+,\S+\s*", t):
        return t.split(",")
    if re.match(r"^\s*[NSEWOnsewo]\s*[-+]?\d", t):
        p = [x for x in re.split(r"\s+(?=[NSEWOnsewo]\s*[-+]?\d)", t) if x.strip()]
        if len(p) == 2:
            return p
    p = [x for x in re.split(r"(?<=[NSEWO])\s+(?=[-+]?\d)|(?<=\d[sS])\s+(?=[-+]?\d)", t) if x.strip()]
    if len(p) == 2:
        return p
    toks = re.findall(r"[-+]?\d+(?:[.,]\d+)?", t)
    if len(toks) == 2:
        return toks
    if len(toks) == 4:
        return [" ".join(toks[:2]), " ".join(toks[2:])]
    if len(toks) == 6:
        return [" ".join(toks[:3]), " ".join(toks[3:])]
    return None


# ======================================================================
# CLASIFICACIÓN DE VALORES POR RANGO (San Juan)
# ======================================================================
# rol: ('lat'|'lon'|'pos_e'|'utm_e'|'norte', valor, sin_signo, tipo 'num'|'txt')
def _rol_angulo(a, positivo, tipo):
    if LAT_ABS[0] <= a <= LAT_ABS[1]:
        return ("lat", -a, positivo, tipo)
    if LON_ABS[0] <= a <= LON_ABS[1]:
        return ("lon", -a, positivo, tipo)
    return None


def rol_celda(v):
    if es_vacio(v) or isinstance(v, bool):
        return None
    n = parse_numero(v)
    if n is not None:
        if abs(n) < 1e5:
            if abs(n - round(n)) < 1e-9:      # enteros chicos: IDs, grados sueltos, etc.
                return None
            return _rol_angulo(abs(n), n > 0, "num")
        if POS_E[0] <= n <= POS_E[1]:
            return ("pos_e", n, False, "num")
        if UTM_E[0] <= n <= UTM_E[1]:
            return ("utm_e", n, False, "num")
        if NORTE[0] <= n <= NORTE[1]:
            return ("norte", n, False, "num")
        return None
    if isinstance(v, str):
        val, _ = parse_angulo_texto(v)
        if val is not None:
            return _rol_angulo(abs(val), val > 0, "txt")
    return None


def rol_par(v):
    """Celda de texto con las dos coordenadas -> (rol_a, rol_b) o None."""
    if not isinstance(v, str) or es_vacio(v) or parse_numero(v) is not None:
        return None
    t = re.sub(r"^\s*(geo:|coordenadas?:?|coords?:?)\s*", "", v.strip(), flags=re.I).strip("()[]{} ")
    m = re.match(r"^\s*POINT\s*Z?\s*\(\s*([-+\d.eE]+)\s+([-+\d.eE]+)", t, re.I)
    if m:
        a, b = float(m.group(1)), float(m.group(2))
        ra, rb = rol_celda(a), rol_celda(b)
        return (ra, rb) if ra and rb else None
    crudos = [parse_numero(x.rstrip(".,")) for x in re.findall(r"[-+]?\d[\d.,]*", t)]
    grandes = [x for x in crudos if x is not None and abs(x) >= 1e5]
    if len(grandes) == 2:
        ra, rb = rol_celda(grandes[0]), rol_celda(grandes[1])
        return (ra, rb) if ra and rb else None
    t = re.sub(r"\b(lat\w*|lon\w*|lng)\s*[:=]?", " ", t, flags=re.I)
    partes = _dividir_par(t)
    if not partes or len(partes) != 2:
        return None
    ra, rb = rol_celda(partes[0]), rol_celda(partes[1])
    if ra is None or rb is None:
        return None
    return ra, rb


def resolver(ra, rb, extra=""):
    """Dos roles -> resultado (sin convertir) o None."""
    if not ra or not rb or ra[0] == rb[0]:
        return None
    roles = {ra[0]: ra, rb[0]: rb}
    claves = set(roles)
    sin_signo = " (sin signo: se asume S/O)" if any(r[2] for r in (ra, rb)) and "lat" in claves else ""
    if claves == {"lat", "lon"}:
        tipo = "WGS84 decimal" if ra[3] == "num" and rb[3] == "num" else "GMS"
        return dict(tipo="geo", lat=roles["lat"][1], lon=roles["lon"][1], sis=tipo + extra + sin_signo)
    if claves == {"pos_e", "norte"}:
        e, n = roles["pos_e"][1], roles["norte"][1]
        return dict(tipo="pos", e=e, n=n, fj=int(e // 1_000_000), sis=f"POSGAR 2007 F{int(e // 1_000_000)}" + extra)
    if claves == {"utm_e", "norte"}:
        return dict(tipo="utm", e=roles["utm_e"][1], n=roles["norte"][1], sis="UTM 19S" + extra)
    return None


# ======================================================================
# DETECCIÓN DE COLUMNAS
# ======================================================================
def roles_por_columna(filas, ncol):
    info = {}
    for c in range(ncol):
        cnt, total = Counter(), 0
        for f in filas:
            v = f[c] if c < len(f) else None
            if es_vacio(v):
                continue
            total += 1
            r = rol_celda(v)
            if r:
                cnt[(r[0], r[3])] += 1
        if cnt:
            (rol, tipo), k = cnt.most_common(1)[0]
            if k / total >= 0.6:
                info[c] = (rol, tipo)
    return info


def _emparejar(A, B):
    B = list(B)
    for a in sorted(A):
        if not B:
            break
        B.sort(key=lambda b: abs(a - b))
        yield a, B.pop(0)


def detectar_triples(filas, ncol):
    """Tres columnas consecutivas grados|minutos|segundos (lat y lon)."""
    encontrados = []
    for c in range(ncol - 2):
        d, m, s, con_dato = [], [], [], 0
        for f in filas:
            trio = [f[c + k] if c + k < len(f) else None for k in range(3)]
            if all(es_vacio(x) for x in trio):
                continue
            con_dato += 1
            p = [parse_numero(x) for x in trio]
            if None in p:
                continue                       # encabezados u otras filas de texto
            d.append(p[0]); m.append(p[1]); s.append(p[2])
        if not d or len(d) / con_dato < 0.6:
            continue
        if not all(abs(x - round(x)) < 1e-9 for x in d + m):
            continue
        if not all(0 <= x < 60 for x in m) or not all(0 <= x < 60 for x in s):
            continue
        ad = [abs(x) for x in d]
        if all(28 <= x <= 33 for x in ad):
            encontrados.append(("lat", c))
        elif all(66 <= x <= 70 for x in ad):
            encontrados.append(("lon", c))
    lat = next((c for r, c in encontrados if r == "lat"), None)
    lon = next((c for r, c in encontrados if r == "lon"), None)
    return (lat, lon) if lat is not None and lon is not None else None


def armar_candidatos(filas, ncol):
    """Lista de candidatos ordenados por prioridad."""
    info = roles_por_columna(filas, ncol)
    cand = []
    for tipo, prio in (("num", 0), ("txt", 3)):
        lat = [c for c, (r, t) in info.items() if r == "lat" and t == tipo]
        lon = [c for c, (r, t) in info.items() if r == "lon" and t == tipo]
        for a, b in _emparejar(lat, lon):
            cand.append(dict(kind="par_cols", a=a, b=b, prio=prio))
    norte = [c for c, (r, _) in info.items() if r == "norte"]
    for rol, prio in (("pos_e", 1), ("utm_e", 2)):
        es = [c for c, (r, _) in info.items() if r == rol]
        for a, b in _emparejar(es, norte):
            norte = [x for x in norte if x != b]
            cand.append(dict(kind="par_cols", a=a, b=b, prio=prio))
    for c in range(ncol):
        if c in info:
            continue
        vals = [f[c] for f in filas if c < len(f) and isinstance(f[c], str) and not es_vacio(f[c])]
        if vals and sum(rol_par(v) is not None for v in vals[:80]) / len(vals[:80]) >= 0.6:
            cand.append(dict(kind="par_celda", a=c, prio=4))
    tri = detectar_triples(filas, ncol)
    if tri:
        cand.append(dict(kind="triples", lat=tri[0], lon=tri[1], prio=5))
    cand.sort(key=lambda x: x["prio"])
    return cand


def evaluar_fila(f, cand):
    g = lambda c: f[c] if c < len(f) else None
    for k in cand:
        if k["kind"] == "par_cols":
            res = resolver(rol_celda(g(k["a"])), rol_celda(g(k["b"])))
        elif k["kind"] == "par_celda":
            p = rol_par(g(k["a"]))
            res = resolver(*p, extra=" (par en una celda)") if p else None
        else:
            res = None
            try:
                vals = []
                for eje in ("lat", "lon"):
                    c = k[eje]
                    d, m, s = parse_numero(g(c)), parse_numero(g(c + 1)), parse_numero(g(c + 2))
                    if d is None or m is None or s is None:
                        raise ValueError
                    vals.append(abs(d) + m / 60 + s / 3600)
                if LAT_ABS[0] <= vals[0] <= LAT_ABS[1] and LON_ABS[0] <= vals[1] <= LON_ABS[1]:
                    res = dict(tipo="geo", lat=-vals[0], lon=-vals[1], sis="GMS (columnas)")
            except ValueError:
                res = None
        if res:
            return res
    return None


def a_wgs84(res):
    """Devuelve (lat, lon) WGS84 o None si cae fuera de San Juan."""
    if res["tipo"] == "geo":
        lat, lon = res["lat"], res["lon"]
    elif res["tipo"] == "pos":
        if res["fj"] not in T_DE_POS:
            return None
        lon, lat = T_DE_POS[res["fj"]].transform(res["e"], res["n"])
    else:
        lon, lat = T_DE_UTM.transform(res["e"], res["n"])
    if not (-33.7 <= lat <= -27.4 and -71.1 <= lon <= -65.7):
        return None
    return lat, lon


def a_gms(val, es_lat):
    hem = ("N" if val >= 0 else "S") if es_lat else ("E" if val >= 0 else "W")
    a = abs(val)
    d = int(a)
    m = int((a - d) * 60)
    s = round((a - d - m / 60) * 3600, 3)
    if s >= 60:
        s -= 60
        m += 1
    if m >= 60:
        m -= 60
        d += 1
    return f"{d}°{m:02d}'{s:06.3f}\"{hem}"


# ======================================================================
# PROCESO POR HOJA / ARCHIVO
# ======================================================================
def procesar_hoja(ws, filas):
    ncol = max((len(f) for f in filas), default=0)
    cand = armar_candidatos(filas, ncol)
    if not cand:
        print(f"   [{ws.title}] no se encontraron coordenadas de San Juan.")
        return 0

    resultados, sin_reconocer = {}, []
    columnas_coord = {c for k in cand for c in (k.get("a"), k.get("b"), k.get("lat"), k.get("lon")) if c is not None}
    primera = None
    for i, f in enumerate(filas):
        res = evaluar_fila(f, cand)
        wgs = a_wgs84(res) if res else None
        if res and wgs:
            resultados[i] = (res, wgs)
            if primera is None:
                primera = i
        elif primera is not None and any(not es_vacio(f[c]) for c in columnas_coord if c < len(f)):
            sin_reconocer.append(i + 1)
    if not resultados:
        print(f"   [{ws.title}] los valores no caen dentro de San Juan.")
        return 0

    # faja POSGAR de la hoja (predominante)
    cuenta = Counter(faja_de_lon(w[1]) for _, w in resultados.values())
    faja = cuenta.most_common(1)[0][0]
    if len(cuenta) > 1:
        print(f"   [!] Puntos en fajas {dict(cuenta)}: se usa la faja {faja} para toda la hoja.")

    # encabezados: fila inmediatamente anterior al primer dato
    desplazo = 0
    start_col = ws.max_column + 1
    if primera == 0:
        ws.insert_rows(1)
        desplazo = 1
        fila_h = 1
    else:
        fila_h = primera          # (primera es índice 0-based del primer dato → fila Excel primera+1; encabezado = primera)
    encabezados = ["LAT_WGS84", "LON_WGS84", f"ESTE_POSGAR07_F{faja}", f"NORTE_POSGAR07_F{faja}",
                   "ESTE_UTM19S", "NORTE_UTM19S", "LAT_GMS", "LON_GMS", "SISTEMA_ORIGEN"]
    relleno = PatternFill("solid", start_color="1F4E78")
    for j, h in enumerate(encabezados):
        c = ws.cell(fila_h, start_col + j, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = relleno
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for i, (res, (lat, lon)) in resultados.items():
        r = i + 1 + desplazo
        if res["tipo"] == "pos" and res["fj"] == faja:
            pe, pn = res["e"], res["n"]
        else:
            pe, pn = T_A_POS[faja].transform(lon, lat)
        if res["tipo"] == "utm":
            ue, un = res["e"], res["n"]
        else:
            ue, un = T_A_UTM.transform(lon, lat)
        vals = [lat, lon, round(pe, 3), round(pn, 3), round(ue, 3), round(un, 3),
                a_gms(lat, True), a_gms(lon, False), res["sis"]]
        for j, v in enumerate(vals):
            c = ws.cell(r, start_col + j, v)
            if j < 2:
                c.number_format = "0.00000000"
            elif 2 <= j < 6:
                c.number_format = "#,##0.000"
    for j, h in enumerate(encabezados):
        ws.column_dimensions[get_column_letter(start_col + j)].width = max(len(h) + 2, 16 if j < 8 else 34)

    print(f"   [{ws.title}] {len(resultados)} puntos convertidos (POSGAR faja {faja}).")
    orig = Counter(r["sis"] for r, _ in resultados.values())
    for s, n in orig.most_common():
        print(f"       · {s}: {n}")
    if sin_reconocer:
        muestra = ", ".join(map(str, sin_reconocer[:10])) + (" …" if len(sin_reconocer) > 10 else "")
        print(f"   [!] {len(sin_reconocer)} fila(s) con datos no reconocidos o fuera de San Juan (filas: {muestra}).")
    return len(resultados)


def procesar_archivo(ruta: Path):
    es_xlsm = ruta.suffix.lower() == ".xlsm"
    wb = load_workbook(ruta, keep_vba=es_xlsm)
    wbv = load_workbook(ruta, data_only=True)
    total = 0
    for ws in wb.worksheets:
        filas = [list(r) for r in wbv[ws.title].iter_rows(values_only=True)]
        total += procesar_hoja(ws, filas)
    if not total:
        print("   (sin cambios: no se genera archivo)")
        return
    salida = ruta.with_name(f"{ruta.stem}_convertido{ruta.suffix}")
    try:
        wb.save(salida)
    except PermissionError:
        print(f"   [ERROR] No se pudo guardar {salida.name}: cerralo en Excel y volvé a ejecutar.")
        return
    print(f"   [OK] Guardado: {salida.name}")


def main():
    carpeta = Path(__file__).resolve().parent
    archivos = sorted(p for p in carpeta.iterdir()
                      if p.is_file() and p.suffix.lower() in (".xlsx", ".xlsm")
                      and not p.name.startswith("~$") and not p.stem.endswith("_convertido"))
    print(f"Carpeta: {carpeta}\nArchivos encontrados: {len(archivos)}")
    for a in archivos:
        print(f"\n▶ {a.name}")
        try:
            procesar_archivo(a)
        except Exception as e:
            print(f"   [ERROR] {e}")
    print("\n✔ Proceso completo")


if __name__ == "__main__":
    main()
