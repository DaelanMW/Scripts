import re
import sys
import zipfile
import xml.etree.ElementTree as ET
from collections import Counter
from pathlib import Path

try:
    from pyproj import Transformer
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError as e:
    sys.exit(f"Falta una librería ({e.name}). Instalá con:  pip install pyproj openpyxl")

# POSGAR 2007 / Argentina faja 1..7  ->  EPSG 5343..5349 (meridiano central -72 + 3*(n-1))
EPSG_POSGAR07_FAJA = {n: 5342 + n for n in range(1, 8)}


# ----------------------------------------------------------------- selección
def buscar_archivos(carpeta: Path):
    return sorted(
        [p for p in carpeta.iterdir() if p.is_file() and p.suffix.lower() in (".kml", ".kmz")],
        key=lambda p: p.name.lower(),
    )


def elegir_archivos(archivos):
    if len(archivos) == 1:
        return archivos

    nombres = [a.name for a in archivos]
    opciones = nombres + ["» Todos los archivos «"]

    try:
        import tkinter as tk
        from tkinter import ttk

        resultado = {}
        root = tk.Tk()
        root.title("KML/KMZ a Excel")
        root.resizable(False, False)
        ttk.Label(root, text="Elegí el archivo a transformar:").pack(padx=15, pady=(15, 5))
        combo = ttk.Combobox(root, values=opciones, state="readonly", width=60)
        combo.current(0)
        combo.pack(padx=15, pady=5)

        def aceptar():
            resultado["i"] = combo.current()
            root.destroy()

        ttk.Button(root, text="Aceptar", command=aceptar).pack(pady=(5, 15))
        root.protocol("WM_DELETE_WINDOW", root.destroy)
        root.mainloop()
        if "i" not in resultado:
            sys.exit("Cancelado.")
        i = resultado["i"]
    except Exception:
        # Sin entorno gráfico: menú por consola
        print("Archivos encontrados:")
        for n, o in enumerate(opciones, 1):
            print(f"  {n}. {o}")
        while True:
            s = input("Elegí un número: ").strip()
            if s.isdigit() and 1 <= int(s) <= len(opciones):
                i = int(s) - 1
                break

    return archivos if i == len(archivos) else [archivos[i]]


# ----------------------------------------------------------------- lectura
def _local(tag):
    return tag.rsplit("}", 1)[-1]


def _texto(el, nombre):
    for c in el:
        if _local(c.tag) == nombre:
            return (c.text or "").strip()
    return ""


def leer_kml_bytes(ruta: Path) -> bytes:
    if ruta.suffix.lower() == ".kmz":
        with zipfile.ZipFile(ruta) as z:
            kmls = [n for n in z.namelist() if n.lower().endswith(".kml")]
            if not kmls:
                raise ValueError("El KMZ no contiene ningún .kml")
            principal = "doc.kml" if "doc.kml" in kmls else kmls[0]
            return z.read(principal)
    return ruta.read_bytes()


def extraer_puntos(ruta: Path):
    """Devuelve lista de dicts con nombre, descripción, carpeta, lon, lat, alt y datos extendidos."""
    raiz = ET.fromstring(leer_kml_bytes(ruta))
    puntos = []

    def procesar_placemark(pm, carpeta):
        nombre = _texto(pm, "name")
        desc = re.sub(r"<[^>]+>", " ", _texto(pm, "description"))
        desc = re.sub(r"\s+", " ", desc).strip()

        extra = {}
        for d in pm.iter():
            t = _local(d.tag)
            if t == "Data":
                v = next((x.text for x in d if _local(x.tag) == "value"), "")
                extra[d.get("name", "")] = (v or "").strip()
            elif t == "SimpleData":
                extra[d.get("name", "")] = (d.text or "").strip()

        for p in pm.iter():
            if _local(p.tag) != "Point":
                continue
            for c in p.iter():
                if _local(c.tag) == "coordinates" and c.text and c.text.strip():
                    partes = c.text.strip().split()[0].split(",")
                    lon, lat = float(partes[0]), float(partes[1])
                    alt = float(partes[2]) if len(partes) > 2 else None
                    puntos.append(
                        dict(nombre=nombre, descripcion=desc, carpeta=carpeta,
                             lon=lon, lat=lat, alt=alt, extra=extra)
                    )

    def recorrer(el, ruta_carpeta):
        for hijo in el:
            t = _local(hijo.tag)
            if t in ("Folder", "Document"):
                n = _texto(hijo, "name")
                recorrer(hijo, ruta_carpeta + ([n] if n else []))
            elif t == "Placemark":
                procesar_placemark(hijo, " / ".join(ruta_carpeta))

    recorrer(raiz, [])
    return puntos


# ----------------------------------------------------------------- conversión
def a_dms(valor, es_lat):
    hemi = ("N" if valor >= 0 else "S") if es_lat else ("E" if valor >= 0 else "W")
    v = abs(valor)
    g = int(v)
    m = int((v - g) * 60)
    s = round((v - g - m / 60) * 3600, 3)
    if s >= 60:
        s -= 60
        m += 1
    if m >= 60:
        m -= 60
        g += 1
    return f"{g}°{m:02d}'{s:06.3f}\"{hemi}"


def faja_posgar(lon):
    n = round((lon + 72) / 3) + 1
    return min(max(n, 1), 7)


# ----------------------------------------------------------------- exportación
def exportar(ruta: Path):
    puntos = extraer_puntos(ruta)
    if not puntos:
        print(f"[!] {ruta.name}: no se encontraron puntos (Point).")
        return None

    # Faja predominante del archivo (todos los puntos se proyectan en la misma faja)
    fajas = Counter(faja_posgar(p["lon"]) for p in puntos)
    faja = fajas.most_common(1)[0][0]
    if len(fajas) > 1:
        print(f"[!] {ruta.name}: puntos en varias fajas {dict(fajas)}. Se usa la faja {faja} para todos.")

    tr = Transformer.from_crs(4326, EPSG_POSGAR07_FAJA[faja], always_xy=True)

    extra_cols = []
    for p in puntos:
        for k in p["extra"]:
            if k not in extra_cols:
                extra_cols.append(k)

    encabezados = [
        "N°", "Nombre", "Descripción", "Carpeta",
        "Latitud WGS84 (°)", "Longitud WGS84 (°)", "Altitud (m)",
        f"Este POSGAR 2007 F{faja} (m)", f"Norte POSGAR 2007 F{faja} (m)",
        "Latitud WGS84 (GMS)", "Longitud WGS84 (GMS)",
    ] + extra_cols

    wb = Workbook()
    ws = wb.active
    ws.title = "Puntos"
    ws.append(encabezados)

    for i, p in enumerate(puntos, 1):
        este, norte = tr.transform(p["lon"], p["lat"])
        ws.append(
            [i, p["nombre"], p["descripcion"], p["carpeta"],
             p["lat"], p["lon"], p["alt"],
             round(este, 3), round(norte, 3),
             a_dms(p["lat"], True), a_dms(p["lon"], False)]
            + [p["extra"].get(k, "") for k in extra_cols]
        )

    # Formato
    relleno = PatternFill("solid", start_color="1F4E78")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = relleno
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for fila in ws.iter_rows(min_row=2):
        fila[4].number_format = "0.00000000"
        fila[5].number_format = "0.00000000"
        fila[6].number_format = "0.00"
        fila[7].number_format = "#,##0.000"
        fila[8].number_format = "#,##0.000"
    for j, col in enumerate(ws.columns, 1):
        ancho = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(j)].width = min(max(ancho + 2, 10), 50)
    ws.freeze_panes = "A2"

    salida = ruta.with_suffix(".xlsx")
    try:
        wb.save(salida)
    except PermissionError:
        sys.exit(f"No se pudo guardar {salida.name}: cerralo en Excel y volvé a ejecutar.")
    print(f"[OK] {ruta.name} -> {salida.name}  ({len(puntos)} puntos, faja {faja})")
    return salida


# ----------------------------------------------------------------- main
def main():
    # Carpeta donde está guardado el script (no el directorio de trabajo de la terminal)
    carpeta = Path(__file__).resolve().parent
    archivos = buscar_archivos(carpeta)
    if not archivos:
        sys.exit(f"No hay archivos .kml ni .kmz en {carpeta}")
    print(f"Carpeta: {carpeta}")

    for a in elegir_archivos(archivos):
        try:
            exportar(a)
        except Exception as e:
            print(f"[ERROR] {a.name}: {e}")


if __name__ == "__main__":
    main()