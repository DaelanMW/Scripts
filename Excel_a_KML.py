#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# Excel -> KML / KMZ  (SAN JUAN, ARGENTINA)
# =========================================
# Convierte los Excel (.xlsx .xlsm) y .csv de la carpeta donde está este script en
# KML y KMZ para Google Earth.  Se puede elegir el archivo (o "Todos") en un menú.
#
# COORDENADAS: se reconocen por los rangos de San Juan, sin importar el nombre ni el
# orden de las columnas (lat/lon, lon/lat, x/y, este/norte, norte/este...):
#     Latitud 28-33 S | Longitud 66-71 O | POSGAR 2007 (fajas 2 y 3) | UTM 19S
# Formatos: decimales (punto o coma), GMS en una celda (31°38'12.5"S, S 31 38 12.5,
# 31d38m12s, 313812S...), GMS en 3 columnas (grados|min|seg), par en una celda
# ("-31.5, -69.4"), con o sin signo, con o sin encabezados, en todas las hojas.
#
# TAMBIÉN RECONOCE (por el nombre del encabezado):
#     Nombre del punto  : Name, Nombre, Punto, ID, Código, Muestra, Estación, Sitio...
#     Descripción       : Descripción, Observaciones, Comentarios, Notas...
#     Grupo (carpetas)  : Carpeta, Grupo, Categoría, Tipo, Capa, Clase...
#     Altura            : Cota, Altura, Elevación, Z, msnm...
# Todas las demás columnas se guardan como atributos y se ven en el globo de
# información de Google Earth, junto con las coordenadas en WGS84, GMS, POSGAR y UTM.
#
# Dependencias:  pip install openpyxl pyproj
import argparse
import csv
import datetime
import html
import math
import re
import subprocess
import sys
import unicodedata
import zipfile
from collections import Counter
from pathlib import Path
from xml.sax.saxutils import escape


def _instalar():
    import importlib
    faltan = []
    for mod, pkg in (("openpyxl", "openpyxl"), ("pyproj", "pyproj")):
        try:
            importlib.import_module(mod)
        except ImportError:
            faltan.append(pkg)
    if faltan:
        print("Instalando:", ", ".join(faltan))
        if subprocess.call([sys.executable, "-m", "pip", "install", *faltan]) != 0:
            sys.exit(f'No se pudo instalar. Ejecutá:\n    "{sys.executable}" -m pip install {" ".join(faltan)}')


_instalar()

from openpyxl import load_workbook
from pyproj import Transformer

# ======================================================================
# CONSTANTES DE SAN JUAN
# ======================================================================
LAT_ABS = (27.5, 33.6)          # |latitud|  (°S)
LON_ABS = (65.8, 71.0)          # |longitud| (°O)
POS_E = (2_200_000, 3_700_000)  # POSGAR 2007 fajas 2 y 3
UTM_E = (250_000, 850_000)      # UTM huso 19S
NORTE = (6_200_000, 7_000_000)  # vale para POSGAR y UTM sur
EPSG_POSGAR = {2: 5344, 3: 5345}
EPSG_UTM19S = 32719

T_A_POS = {f: Transformer.from_crs(4326, e, always_xy=True) for f, e in EPSG_POSGAR.items()}
T_DE_POS = {f: Transformer.from_crs(e, 4326, always_xy=True) for f, e in EPSG_POSGAR.items()}
T_A_UTM = Transformer.from_crs(4326, EPSG_UTM19S, always_xy=True)
T_DE_UTM = Transformer.from_crs(EPSG_UTM19S, 4326, always_xy=True)

VACIOS = {"", "nan", "none", "null", "nd", "s/d", "sd", "n/a", "na", "-", "--", "#n/a", "#value!", "s/n"}


def faja_de_lon(lon):
    # San Juan usa fajas 2 (hasta 67°30'O) y 3.
    return 2 if lon < -67.5 else 3


# ======================================================================
# LECTURA DE VALORES
# ======================================================================
def es_vacio(v):
    if v is None:
        return True
    if isinstance(v, float) and math.isnan(v):
        return True
    return isinstance(v, str) and v.strip().lower() in VACIOS


_RE_NUM = re.compile(r"^[-+]?(\d+\.?\d*|\.\d+)([eE][-+]?\d+)?$")


def parse_numero(v):
    if es_vacio(v) or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v) if math.isfinite(v) else None
    if not isinstance(v, str):
        return None
    s = v.strip().replace("\u00a0", " ").replace("\u202f", " ")
    s = s.replace("−", "-").replace("–", "-").replace("—", "-")
    s = re.sub(r"\s*(mts?|metros?|m)\.?$", "", s, flags=re.I)
    if re.fullmatch(r"[-+]?\d{1,3}( \d{3})+([.,]\d+)?", s):
        s = s.replace(" ", "")
    if " " in s:
        return None
    if "," in s and "." in s:
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif "," in s:
        if re.fullmatch(r"[-+]?\d{1,3}(,\d{3}){2,}", s):
            s = s.replace(",", "")
        else:
            s = s.replace(",", ".")
    elif s.count(".") > 1:
        if re.fullmatch(r"[-+]?\d{1,3}(\.\d{3})+", s):
            s = s.replace(".", "")
        else:
            return None
    if _RE_NUM.match(s):
        x = float(s)
        return x if math.isfinite(x) else None
    return None


def _compacto(x):
    # DDMMSS(.s) / DDDMMSS(.s) / DDMM(.m) escrito como un único número.
    a = abs(x)
    ent = int(a)
    frac = a - ent
    n = len(str(ent))
    if n in (6, 7):
        ss, mm, dd = ent % 100 + frac, (ent // 100) % 100, ent // 10000
    elif n in (4, 5):
        mm, ss, dd = ent % 100 + frac, 0, ent // 100
    else:
        return None
    if mm >= 60 or ss >= 60 or dd > 180:
        return None
    return dd + mm / 60 + ss / 3600


def _armar(partes, hem):
    if hem:
        hem = {"O": "W"}.get(hem.upper(), hem.upper())
    if not 1 <= len(partes) <= 3:
        return None, None
    try:
        nums = [float(p) for p in partes]
    except ValueError:
        return None, None
    neg = partes[0].startswith("-") or math.copysign(1, nums[0]) < 0 or hem in ("S", "W")
    d = abs(nums[0])
    if len(nums) == 1:
        if d >= 1000:
            val = _compacto(d)
            if val is None:
                return None, None
        elif d > 180:
            return None, None
        else:
            val = d
    else:
        if d != int(d):
            return None, None
        m = nums[1]
        if m < 0 or m >= 60:
            return None, None
        val = d + m / 60
        if len(nums) == 3:
            s = nums[2]
            if m != int(m) or s < 0 or s >= 60:
                return None, None
            val += s / 3600
    return (-val if neg else val), hem


_PALABRAS_HEM = [(r"\bsur\b|\bsouth\b", "S"), (r"\bnorte\b|\bnorth\b", "N"),
                 (r"\boeste\b|\bwest\b", "W"), (r"\beste\b|\beast\b", "E")]
_RE_LETRAS_DMS = re.compile(
    r"\s*([NSEWOnsewo])?[\s:=]*([-+]?\d+(?:[.,]\d+)?)\s*[dD]\s*"
    r"(?:(\d+(?:[.,]\d+)?)\s*[mM]\s*(?:(\d+(?:[.,]\d+)?)\s*[sS]?)?)?\s*([NSEWO])?\s*")


def parse_angulo_texto(s):
    # Texto en GMS (o decimal con letras) -> (grados decimales, hemisferio|None).
    t = s.strip()
    if not t:
        return None, None
    t = t.replace("−", "-").replace("–", "-").replace("—", "-")
    t = re.sub(r"[º˚°]", "°", t)
    t = re.sub(r"[′’´`‘]", "'", t)
    t = re.sub(r"[″”“]", '"', t)
    t = t.replace("''", '"')
    hem = None
    for pat, h in _PALABRAS_HEM:
        if re.search(pat, t, re.I):
            hem = h
            t = re.sub(pat, " ", t, flags=re.I)
    m = _RE_LETRAS_DMS.fullmatch(t)
    if m:
        partes = [p.replace(",", ".") for p in (m.group(2), m.group(3), m.group(4)) if p]
        return _armar(partes, m.group(1) or m.group(5) or hem)
    m = re.match(r"^\s*([NSEWOnsewo])[\s:=.]*(?=[-+]?\d)", t)
    if m:
        hem = hem or m.group(1).upper()
        t = t[m.end():]
    else:
        m = re.search(r"(?<=[\d'\"°.\s])\s*([NSEWOnsewo])\s*$", t)
        if m:
            hem = hem or m.group(1).upper()
            t = t[:m.start()]
    t = re.sub(r"\b(grados?|degrees?|deg|minutos?|mins?|minutes?|segundos?|segs?|secs?|seconds?|"
               r"latitud|latitude|lat|longitud|longitude|long|lon|lng|coord\w*)\b", " ", t, flags=re.I)
    if re.search(r"[A-Za-z]", t):
        return None, None
    t = re.sub(r"[°'\"]", " ", t).replace(":", " ")
    t = re.sub(r"(?<=\d)\s*-\s*(?=\d)", " ", t)
    if t.count(",") >= 2 and "." not in t:
        t = t.replace(",", " ")
    else:
        t = t.replace(",", ".")
    patron = r"[-+]?\d+(?:\.\d+)?"
    partes = re.findall(patron, t)
    if re.sub(patron, "", t).strip(" \t.;"):
        return None, None
    return _armar(partes, hem)


def _dividir_par(t):
    partes = [p for p in re.split(r"\s*[;|\t\n]\s*|\s+/\s+|\s+y\s+", t) if p.strip()]
    if len(partes) == 2:
        return partes
    partes = [p for p in re.split(r",\s+", t) if p.strip()]
    if len(partes) == 2:
        return partes
    if t.count(",") == 1 and re.fullmatch(r"\s*\S+,\S+\s*", t):
        return t.split(",")
    if re.match(r"^\s*[NSEWOnsewo]\s*[-+]?\d", t):
        p = [x for x in re.split(r"\s+(?=[NSEWOnsewo]\s*[-+]?\d)", t) if x.strip()]
        if len(p) == 2:
            return p
    p = [x for x in re.split(r"(?<=[NSEWO])\s+(?=[-+]?\d)|(?<=\d[sS])\s+(?=[-+]?\d)", t) if x.strip()]
    if len(p) == 2:
        return p
    toks = re.findall(r"[-+]?\d+(?:[.,]\d+)?", t)
    if len(toks) == 2:
        return toks
    if len(toks) == 4:
        return [" ".join(toks[:2]), " ".join(toks[2:])]
    if len(toks) == 6:
        return [" ".join(toks[:3]), " ".join(toks[3:])]
    return None


# ======================================================================
# CLASIFICACIÓN DE VALORES POR RANGO (San Juan)
# ======================================================================
# rol: ('lat'|'lon'|'pos_e'|'utm_e'|'norte', valor, sin_signo, tipo 'num'|'txt')
def _rol_angulo(a, positivo, tipo):
    if LAT_ABS[0] <= a <= LAT_ABS[1]:
        return ("lat", -a, positivo, tipo)
    if LON_ABS[0] <= a <= LON_ABS[1]:
        return ("lon", -a, positivo, tipo)
    return None


def rol_celda(v):
    if es_vacio(v) or isinstance(v, bool):
        return None
    n = parse_numero(v)
    if n is not None:
        if abs(n) < 1e5:
            if abs(n - round(n)) < 1e-9:      # enteros chicos: IDs, grados sueltos, etc.
                return None
            return _rol_angulo(abs(n), n > 0, "num")
        if POS_E[0] <= n <= POS_E[1]:
            return ("pos_e", n, False, "num")
        if UTM_E[0] <= n <= UTM_E[1]:
            return ("utm_e", n, False, "num")
        if NORTE[0] <= n <= NORTE[1]:
            return ("norte", n, False, "num")
        return None
    if isinstance(v, str):
        val, _ = parse_angulo_texto(v)
        if val is not None:
            return _rol_angulo(abs(val), val > 0, "txt")
    return None


def rol_par(v):
    # Celda de texto con las dos coordenadas -> (rol_a, rol_b) o None.
    if not isinstance(v, str) or es_vacio(v) or parse_numero(v) is not None:
        return None
    t = re.sub(r"^\s*(geo:|coordenadas?:?|coords?:?)\s*", "", v.strip(), flags=re.I).strip("()[]{} ")
    m = re.match(r"^\s*POINT\s*Z?\s*\(\s*([-+\d.eE]+)\s+([-+\d.eE]+)", t, re.I)
    if m:
        a, b = float(m.group(1)), float(m.group(2))
        ra, rb = rol_celda(a), rol_celda(b)
        return (ra, rb) if ra and rb else None
    crudos = [parse_numero(x.rstrip(".,")) for x in re.findall(r"[-+]?\d[\d.,]*", t)]
    grandes = [x for x in crudos if x is not None and abs(x) >= 1e5]
    if len(grandes) == 2:
        ra, rb = rol_celda(grandes[0]), rol_celda(grandes[1])
        return (ra, rb) if ra and rb else None
    t = re.sub(r"\b(lat\w*|lon\w*|lng)\s*[:=]?", " ", t, flags=re.I)
    partes = _dividir_par(t)
    if not partes or len(partes) != 2:
        return None
    ra, rb = rol_celda(partes[0]), rol_celda(partes[1])
    if ra is None or rb is None:
        return None
    return ra, rb


def resolver(ra, rb, extra=""):
    # Dos roles -> resultado (sin convertir) o None.
    if not ra or not rb or ra[0] == rb[0]:
        return None
    roles = {ra[0]: ra, rb[0]: rb}
    claves = set(roles)
    sin_signo = " (sin signo: se asume S/O)" if any(r[2] for r in (ra, rb)) and "lat" in claves else ""
    if claves == {"lat", "lon"}:
        tipo = "WGS84 decimal" if ra[3] == "num" and rb[3] == "num" else "GMS"
        return dict(tipo="geo", lat=roles["lat"][1], lon=roles["lon"][1], sis=tipo + extra + sin_signo)
    if claves == {"pos_e", "norte"}:
        e, n = roles["pos_e"][1], roles["norte"][1]
        return dict(tipo="pos", e=e, n=n, fj=int(e // 1_000_000), sis=f"POSGAR 2007 F{int(e // 1_000_000)}" + extra)
    if claves == {"utm_e", "norte"}:
        return dict(tipo="utm", e=roles["utm_e"][1], n=roles["norte"][1], sis="UTM 19S" + extra)
    return None


# ======================================================================
# DETECCIÓN DE COLUMNAS
# ======================================================================
def roles_por_columna(filas, ncol):
    info = {}
    for c in range(ncol):
        cnt, total = Counter(), 0
        for f in filas:
            v = f[c] if c < len(f) else None
            if es_vacio(v):
                continue
            total += 1
            r = rol_celda(v)
            if r:
                cnt[(r[0], r[3])] += 1
        if cnt:
            (rol, tipo), k = cnt.most_common(1)[0]
            if k / total >= 0.6:
                info[c] = (rol, tipo)
    return info


def _emparejar(A, B):
    B = list(B)
    for a in sorted(A):
        if not B:
            break
        B.sort(key=lambda b: abs(a - b))
        yield a, B.pop(0)


def detectar_triples(filas, ncol):
    # Tres columnas consecutivas grados|minutos|segundos (lat y lon).
    encontrados = []
    for c in range(ncol - 2):
        d, m, s, con_dato = [], [], [], 0
        for f in filas:
            trio = [f[c + k] if c + k < len(f) else None for k in range(3)]
            if all(es_vacio(x) for x in trio):
                continue
            con_dato += 1
            p = [parse_numero(x) for x in trio]
            if None in p:
                continue                       # encabezados u otras filas de texto
            d.append(p[0]); m.append(p[1]); s.append(p[2])
        if not d or len(d) / con_dato < 0.6:
            continue
        if not all(abs(x - round(x)) < 1e-9 for x in d + m):
            continue
        if not all(0 <= x < 60 for x in m) or not all(0 <= x < 60 for x in s):
            continue
        ad = [abs(x) for x in d]
        if all(28 <= x <= 33 for x in ad):
            encontrados.append(("lat", c))
        elif all(66 <= x <= 70 for x in ad):
            encontrados.append(("lon", c))
    lat = next((c for r, c in encontrados if r == "lat"), None)
    lon = next((c for r, c in encontrados if r == "lon"), None)
    return (lat, lon) if lat is not None and lon is not None else None


def armar_candidatos(filas, ncol):
    # Lista de candidatos ordenados por prioridad.
    info = roles_por_columna(filas, ncol)
    cand = []
    for tipo, prio in (("num", 0), ("txt", 3)):
        lat = [c for c, (r, t) in info.items() if r == "lat" and t == tipo]
        lon = [c for c, (r, t) in info.items() if r == "lon" and t == tipo]
        for a, b in _emparejar(lat, lon):
            cand.append(dict(kind="par_cols", a=a, b=b, prio=prio))
    norte = [c for c, (r, _) in info.items() if r == "norte"]
    for rol, prio in (("pos_e", 1), ("utm_e", 2)):
        es = [c for c, (r, _) in info.items() if r == rol]
        for a, b in _emparejar(es, norte):
            norte = [x for x in norte if x != b]
            cand.append(dict(kind="par_cols", a=a, b=b, prio=prio))
    for c in range(ncol):
        if c in info:
            continue
        vals = [f[c] for f in filas if c < len(f) and isinstance(f[c], str) and not es_vacio(f[c])]
        if vals and sum(rol_par(v) is not None for v in vals[:80]) / len(vals[:80]) >= 0.6:
            cand.append(dict(kind="par_celda", a=c, prio=4))
    tri = detectar_triples(filas, ncol)
    if tri:
        cand.append(dict(kind="triples", lat=tri[0], lon=tri[1], prio=5))
    cand.sort(key=lambda x: x["prio"])
    return cand


def evaluar_fila(f, cand):
    g = lambda c: f[c] if c < len(f) else None
    for k in cand:
        if k["kind"] == "par_cols":
            res = resolver(rol_celda(g(k["a"])), rol_celda(g(k["b"])))
        elif k["kind"] == "par_celda":
            p = rol_par(g(k["a"]))
            res = resolver(*p, extra=" (par en una celda)") if p else None
        else:
            res = None
            try:
                vals = []
                for eje in ("lat", "lon"):
                    c = k[eje]
                    d, m, s = parse_numero(g(c)), parse_numero(g(c + 1)), parse_numero(g(c + 2))
                    if d is None or m is None or s is None:
                        raise ValueError
                    vals.append(abs(d) + m / 60 + s / 3600)
                if LAT_ABS[0] <= vals[0] <= LAT_ABS[1] and LON_ABS[0] <= vals[1] <= LON_ABS[1]:
                    res = dict(tipo="geo", lat=-vals[0], lon=-vals[1], sis="GMS (columnas)")
            except ValueError:
                res = None
        if res:
            return res
    return None


def a_wgs84(res):
    # Devuelve (lat, lon) WGS84 o None si cae fuera de San Juan.
    if res["tipo"] == "geo":
        lat, lon = res["lat"], res["lon"]
    elif res["tipo"] == "pos":
        if res["fj"] not in T_DE_POS:
            return None
        lon, lat = T_DE_POS[res["fj"]].transform(res["e"], res["n"])
    else:
        lon, lat = T_DE_UTM.transform(res["e"], res["n"])
    if not (-33.7 <= lat <= -27.4 and -71.1 <= lon <= -65.7):
        return None
    return lat, lon


def a_gms(val, es_lat):
    hem = ("N" if val >= 0 else "S") if es_lat else ("E" if val >= 0 else "W")
    a = abs(val)
    d = int(a)
    m = int((a - d) * 60)
    s = round((a - d - m / 60) * 3600, 3)
    if s >= 60:
        s -= 60
        m += 1
    if m >= 60:
        m -= 60
        d += 1
    return f"{d}°{m:02d}'{s:06.3f}\"{hem}"


# ======================================================================
# NOMBRE / DESCRIPCIÓN / GRUPO / ALTURA (por encabezado)
# ======================================================================
def norm(t):
    t = str(t)
    t = re.sub(r"(?<=[a-z])(?=[A-Z])", " ", t)
    t = "".join(c for c in unicodedata.normalize("NFD", t) if unicodedata.category(c) != "Mn").lower()
    return re.sub(r"[^a-z0-9]+", " ", t).strip()


# En orden de prioridad
_NOMBRE = ["name", "nombre", "nombres", "punto", "pto", "id", "codigo", "cod", "code", "muestra", "estacion",
           "sitio", "etiqueta", "label", "titulo", "identificador", "denominacion", "labor", "sondaje",
           "hito", "vertice", "mojon", "waypoint", "wpt", "ref", "referencia", "numero", "nro", "n"]
_DESC = {"descripcion", "description", "desc", "observaciones", "observacion", "obs", "comentario",
         "comentarios", "notas", "nota", "detalle", "detalles", "comments", "notes", "remarks"}
_GRUPO = {"carpeta", "grupo", "categoria", "tipo", "capa", "layer", "folder", "clase", "clasificacion",
          "group", "category", "type", "class", "unidad", "litologia", "seccion"}
_ALT = {"alt", "altura", "altitud", "altitude", "elevacion", "elevation", "elev", "cota", "z", "msnm", "height"}


def _tokens(h):
    return norm(h).split() if h else []


def buscar_columna(headers, vocabulario, excluidas):
    # vocabulario: lista (prioridad) o set. Devuelve índice de columna o None.
    orden = list(vocabulario) if isinstance(vocabulario, list) else None
    mejor = None
    for c, h in enumerate(headers):
        if c in excluidas or not h:
            continue
        toks = _tokens(h)
        if not toks or len(toks) > 4:
            continue
        for t in toks:
            if orden is not None:
                if t in orden:
                    pri = orden.index(t)
                    if mejor is None or pri < mejor[0]:
                        mejor = (pri, c)
            elif t in vocabulario:
                return c
    return mejor[1] if mejor else None


def texto_celda(v):
    if es_vacio(v):
        return ""
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.isoformat(sep=" ")[:19] if isinstance(v, datetime.datetime) else v.isoformat()
    if isinstance(v, bool):
        return "Sí" if v else "No"
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else format(v, ".12g")
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(v)).strip()


# ======================================================================
# LECTURA DE ARCHIVOS
# ======================================================================
def leer_hojas(ruta: Path):
    # {nombre_hoja: filas}
    if ruta.suffix.lower() == ".csv":
        for enc in ("utf-8-sig", "latin-1"):
            try:
                with open(ruta, newline="", encoding=enc) as f:
                    muestra = f.read(4096)
                    f.seek(0)
                    try:
                        dial = csv.Sniffer().sniff(muestra, delimiters=",;\t|")
                    except csv.Error:
                        dial = csv.excel
                    return {"csv": [row for row in csv.reader(f, dial)]}
            except UnicodeDecodeError:
                continue
        raise ValueError("No se pudo leer el CSV")
    wb = load_workbook(ruta, data_only=True, read_only=True)
    hojas = {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}
    wb.close()
    return hojas


# ======================================================================
# EXTRACCIÓN DE PUNTOS
# ======================================================================
def columnas_de_candidatos(cand):
    cols = set()
    for k in cand:
        for clave in ("a", "b"):
            if k.get(clave) is not None:
                cols.add(k[clave])
        for eje in ("lat", "lon"):
            if k.get(eje) is not None and k["kind"] == "triples":
                cols.update((k[eje], k[eje] + 1, k[eje] + 2))
    return cols


def extraer_hoja(nombre_hoja, filas):
    filas = [list(f) for f in filas]
    ncol = max((len(f) for f in filas), default=0)
    cand = armar_candidatos(filas, ncol)
    if not cand:
        print(f"   [{nombre_hoja}] no se encontraron coordenadas de San Juan.")
        return []
    cols_coord = columnas_de_candidatos(cand)

    hallados, sin_reconocer, primera = [], [], None
    for i, f in enumerate(filas):
        res = evaluar_fila(f, cand)
        wgs = a_wgs84(res) if res else None
        if res and wgs:
            hallados.append((i, res, wgs))
            if primera is None:
                primera = i
        elif primera is not None and any(not es_vacio(f[c]) for c in cols_coord if c < len(f)):
            sin_reconocer.append(i + 1)
    if not hallados:
        print(f"   [{nombre_hoja}] los valores no caen dentro de San Juan.")
        return []

    # encabezados: la fila no vacía más cercana por encima del primer dato
    headers = [""] * ncol
    for r in range(primera - 1, max(primera - 4, -1), -1):
        if r >= 0 and any(not es_vacio(v) for v in filas[r]):
            headers = [("" if es_vacio(v) else texto_celda(v)) for v in filas[r]] + [""] * (ncol - len(filas[r]))
            break

    usadas = set(cols_coord)
    c_desc = buscar_columna(headers, _DESC, usadas)
    if c_desc is not None:
        usadas.add(c_desc)
    c_alt = buscar_columna(headers, _ALT, usadas)
    if c_alt is not None:
        usadas.add(c_alt)
    c_grupo = buscar_columna(headers, _GRUPO, usadas)
    c_nombre = buscar_columna(headers, _NOMBRE, usadas | ({c_grupo} if c_grupo is not None else set()))
    if c_nombre is None:                       # sin encabezado útil: se deduce por el contenido
        def valores(c):
            return [filas[i][c] for i, _, _ in hallados[:20] if c < len(filas[i]) and not es_vacio(filas[i][c])]
        candidatas = [c for c in range(ncol) if c not in usadas and c != c_grupo and valores(c)]
        textos = [c for c in candidatas if all(isinstance(v, str) for v in valores(c))]
        enteros = [c for c in candidatas if all(isinstance(v, int) or (isinstance(v, float) and v.is_integer())
                                                for v in valores(c))]
        # primero texto (P1, M-3, Sitio 4...), después un número entero correlativo (ID)
        for lista in (textos, enteros):
            if lista:
                c_nombre = lista[0]
                break
    if c_grupo is not None:
        distintos = {texto_celda(filas[i][c_grupo]) for i, _, _ in hallados if c_grupo < len(filas[i])}
        distintos.discard("")
        if not 1 <= len(distintos) <= 30:
            c_grupo = None
    if c_nombre is not None:
        usadas.add(c_nombre)

    # nombres únicos de atributos
    nombres_attr = {}
    vistos = Counter()
    for c in range(ncol):
        base = headers[c] or f"Columna {c + 1}"
        vistos[base] += 1
        nombres_attr[c] = base if vistos[base] == 1 else f"{base} ({vistos[base]})"

    puntos = []
    for i, res, (lat, lon) in hallados:
        f = filas[i] + [None] * (ncol - len(filas[i]))
        alt = None
        if c_alt is not None:
            a = parse_numero(f[c_alt])
            if a is not None and -500 <= a <= 7500:
                alt = a
        nombre = texto_celda(f[c_nombre]) if c_nombre is not None else ""
        attrs = [(nombres_attr[c], texto_celda(f[c])) for c in range(ncol)
                 if c not in cols_coord and c not in (c_desc, c_nombre) and texto_celda(f[c])]
        puntos.append(dict(
            fila=i + 1, nombre=nombre or f"Punto {len(puntos) + 1}", lat=lat, lon=lon, alt=alt, res=res,
            desc=texto_celda(f[c_desc]) if c_desc is not None else "",
            grupo=texto_celda(f[c_grupo]) if c_grupo is not None else "", attrs=attrs))

    print(f"   [{nombre_hoja}] {len(puntos)} puntos.")
    usados = {"nombre": c_nombre, "descripción": c_desc, "grupo": c_grupo, "altura": c_alt}
    det = ", ".join(f"{k}='{headers[v] or 'col ' + str(v + 1)}'" for k, v in usados.items() if v is not None)
    if det:
        print(f"       · Columnas reconocidas: {det}")
    if c_nombre is None:
        print("       · Sin columna de nombre: se numeran como 'Punto 1', 'Punto 2'...")
    for s, n in Counter(p["res"]["sis"] for p in puntos).most_common():
        print(f"       · Origen {s}: {n}")
    if sin_reconocer:
        muestra = ", ".join(map(str, sin_reconocer[:10])) + (" …" if len(sin_reconocer) > 10 else "")
        print(f"   [!] {len(sin_reconocer)} fila(s) no reconocidas o fuera de San Juan (filas: {muestra}).")
    return puntos


# ======================================================================
# KML
# ======================================================================
COLORES = ["ff0000ff", "ffff0000", "ff00aa00", "ff00ffff", "ff0080ff", "ffaa00aa", "ffffff00", "ffaa55ff"]


def cdata(s):
    return "<![CDATA[" + s.replace("]]>", "]]]]><![CDATA[>") + "]]>"


def globo(p, faja_pt):
    lat, lon = p["lat"], p["lon"]
    pe, pn = T_A_POS[faja_pt].transform(lon, lat)
    ue, un = T_A_UTM.transform(lon, lat)
    fil = lambda k, v: f"<tr><td><b>{html.escape(k)}</b></td><td>{html.escape(v)}</td></tr>"
    h = ""
    if p["desc"]:
        h += f"<p>{html.escape(p['desc']).replace(chr(10), '<br>')}</p>"
    h += "<table border='1' cellpadding='3' cellspacing='0'>"
    for k, v in p["attrs"]:
        h += fil(k, v)
    h += fil("Lat/Lon WGS84", f"{lat:.7f}, {lon:.7f}")
    h += fil("Lat/Lon GMS", f"{a_gms(lat, True)}  {a_gms(lon, False)}")
    h += fil(f"POSGAR 2007 F{faja_pt}", f"E {pe:,.2f}  N {pn:,.2f}")
    h += fil("UTM 19S", f"E {ue:,.2f}  N {un:,.2f}")
    h += fil("Origen", p["res"]["sis"])
    return h + "</table>"


def placemark(p, estilo):
    faja_pt = faja_de_lon(p["lon"])
    coord = f"{p['lon']:.8f},{p['lat']:.8f}" + (f",{p['alt']:.2f}" if p["alt"] is not None else "")
    ext = "".join(f'<Data name="{escape(k, {chr(34): "&quot;"})}"><value>{escape(v)}</value></Data>'
                  for k, v in p["attrs"])
    return (f"<Placemark><name>{escape(p['nombre'])}</name>"
            f"<description>{cdata(globo(p, faja_pt))}</description>"
            f"{estilo}<ExtendedData>{ext}</ExtendedData>"
            f"<Point><coordinates>{coord}</coordinates></Point></Placemark>\n")


def generar_kml(titulo, hojas_puntos):
    # hojas_puntos: {hoja: [puntos]} -> texto KML.
    grupos_global = []
    for pts in hojas_puntos.values():
        for p in pts:
            if p["grupo"] and p["grupo"] not in grupos_global:
                grupos_global.append(p["grupo"])
    estilo_de = {g: i for i, g in enumerate(grupos_global)}

    out = ['<?xml version="1.0" encoding="UTF-8"?>\n<kml xmlns="http://www.opengis.net/kml/2.2">\n<Document>\n',
           f"<name>{escape(titulo)}</name>\n"]
    for g, i in estilo_de.items():
        out.append(f'<Style id="g{i}"><IconStyle><color>{COLORES[i % len(COLORES)]}</color><scale>1.1</scale>'
                   f'<Icon><href>http://maps.google.com/mapfiles/kml/paddle/wht-blank.png</href></Icon>'
                   f"</IconStyle></Style>\n")
    varias = len(hojas_puntos) > 1
    for hoja, pts in hojas_puntos.items():
        if varias:
            out.append(f"<Folder><name>{escape(hoja)}</name>\n")
        por_grupo = {}
        for p in pts:
            por_grupo.setdefault(p["grupo"], []).append(p)
        for g, lista in por_grupo.items():
            if g:
                out.append(f"<Folder><name>{escape(g)}</name>\n")
            est = f"<styleUrl>#g{estilo_de[g]}</styleUrl>" if g else ""
            out.extend(placemark(p, est) for p in lista)
            if g:
                out.append("</Folder>\n")
        if varias:
            out.append("</Folder>\n")
    out.append("</Document>\n</kml>\n")
    return "".join(out)


def procesar_archivo(ruta: Path):
    hojas = leer_hojas(ruta)
    resultado = {}
    for nombre, filas in hojas.items():
        pts = extraer_hoja(nombre, filas)
        if pts:
            resultado[nombre] = pts
    if not resultado:
        print("   (no se genera KML)")
        return
    kml_path = ruta.with_suffix(".kml")
    kmz_path = ruta.with_suffix(".kmz")
    try:
        kml_path.write_text(generar_kml(ruta.stem, resultado), encoding="utf-8")
        with zipfile.ZipFile(kmz_path, "w", zipfile.ZIP_DEFLATED) as z:
            z.write(kml_path, arcname="doc.kml")
    except PermissionError:
        print("   [ERROR] No se pudo escribir el KML/KMZ (¿abierto en Google Earth?).")
        return
    total = sum(len(v) for v in resultado.values())
    print(f"   [OK] {kml_path.name} y {kmz_path.name}  ({total} puntos)")


# ======================================================================
# MENÚ
# ======================================================================
def elegir(archivos):
    if len(archivos) == 1:
        return archivos
    opciones = [a.name for a in archivos] + ["» Todos los archivos «"]
    try:
        import tkinter as tk
        from tkinter import ttk
        res = {}
        root = tk.Tk()
        root.title("Excel → KML/KMZ")
        root.resizable(False, False)
        ttk.Label(root, text="Elegí el archivo a convertir:").pack(padx=15, pady=(15, 5))
        combo = ttk.Combobox(root, values=opciones, state="readonly", width=60)
        combo.current(0)
        combo.pack(padx=15, pady=5)

        def ok():
            res["i"] = combo.current()
            root.destroy()

        ttk.Button(root, text="Aceptar", command=ok).pack(pady=(5, 15))
        root.mainloop()
        if "i" not in res:
            sys.exit("Cancelado.")
        i = res["i"]
    except Exception as e:
        if isinstance(e, SystemExit):
            raise
        print("Archivos encontrados:")
        for n, o in enumerate(opciones, 1):
            print(f"  {n}. {o}")
        s = input("Número [1]: ").strip()
        i = int(s) - 1 if s.isdigit() and 1 <= int(s) <= len(opciones) else 0
    return archivos if i == len(archivos) else [archivos[i]]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--sin-menu", action="store_true", help="procesa todos los archivos sin abrir ventanas")
    ap.add_argument("--archivo")
    args = ap.parse_args()
    carpeta = Path(__file__).resolve().parent
    if args.archivo:
        archivos = [Path(args.archivo).resolve()]
    else:
        archivos = sorted((p for p in carpeta.iterdir()
                           if p.is_file() and p.suffix.lower() in (".xlsx", ".xlsm", ".csv")
                           and not p.name.startswith("~$")), key=lambda p: p.name.lower())
    if not archivos:
        sys.exit(f"No se encontró ningún Excel ni CSV en {carpeta}")
    seleccion = archivos if args.sin_menu else elegir(archivos)
    print(f"Carpeta: {carpeta}")
    for a in seleccion:
        print(f"\n▶ {a.name}")
        try:
            procesar_archivo(a)
        except Exception as e:
            print(f"   [ERROR] {e}")
    print("\n✔ Proceso completo")


if __name__ == "__main__":
    main()
